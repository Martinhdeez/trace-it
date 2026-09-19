"""The importer preserves history and produces reviewable, repeatable reference data."""

from pathlib import Path

import openpyxl
import pytest

from prepare_batch2 import prepare

MATERIAL = Path(__file__).resolve().parents[1] / ".context/500-sombras-de-alberto"
pytestmark = pytest.mark.skipif(
    not (MATERIAL / "pedidos_nuevos.csv").is_file(), reason="Batch 2 material required"
)


def test_cumulative_import_is_idempotent_and_preserves_original_rows(tmp_path):
    book = MATERIAL / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"
    suppliers, orders = MATERIAL / "proveedores_nuevos.csv", MATERIAL / "pedidos_nuevos.csv"
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    result = prepare(book, suppliers, orders, first)
    assert result["tables"] == {
        "Proveedores": {"added": 4, "rows": 16},
        "Pedidos_2026": {"added": 39, "rows": 555},
    }
    repeated = prepare(first, suppliers, orders, second)
    assert all(t["added"] == 0 for t in repeated["tables"].values())
    original = openpyxl.load_workbook(book)
    updated = openpyxl.load_workbook(first)
    again = openpyxl.load_workbook(second)
    for name in original.sheetnames:
        old = list(original[name].values)
        new = list(updated[name].values)
        assert new[: len(old)] == old
        assert list(again[name].values) == new


def test_conflict_fails_without_writing_a_partial_workbook(tmp_path):
    suppliers = tmp_path / "suppliers.csv"
    suppliers.write_text(
        "ID,Razon Social,NIF,IBAN,Ciudad,Condiciones\n"
        "P001,Changed,B12345678,ES00123456789,Madrid,30 dias\n",
        encoding="utf-8",
    )
    target = tmp_path / "invalid.xlsx"
    with pytest.raises(ValueError, match="Conflicting"):
        prepare(
            MATERIAL / "FINAL_v7_DEFINITIVO_ahorasi.xlsx",
            suppliers,
            MATERIAL / "pedidos_nuevos.csv",
            target,
        )
    assert not target.exists()
