"""The real API: workbook/PDF -> stored symbols -> rules -> JSONL, with no LLM keys."""

import asyncio
import io
import json
import os
import uuid

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import select

from app.core.database import engine, session_factory
from app.core.events import Event
from app.features.ingestion.model import Instance
from app.features.ingestion.runtime import current_service
from app.features.ingestion.service import ExtractionService
from app.features.processes.model import DecisionType, Symbol
from app.features.rules.model import Rule
from app.features.sources.model import Source
from app.features.users.model import User
from app.main import app
from tests.support import pack, rows

from .conftest import VALID, NoOCR, NoVLM, pdf_bytes

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(
        os.getenv("TRACEPAY_TEST_POSTGRES") != "1",
        reason="Requires test PostgreSQL",
    ),
]


def workbook(incomplete=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Proveedores"
    sheet.append(["ID", "Razon Social", "NIF", "IBAN"])
    sheet.append(["P001", "Limpiezas Turia", "B98120774", "ES4414650100951704302211"])
    orders = book.create_sheet("Pedidos_2026")
    orders.append(["Pedido", "ProveedorID", "NIF", "Importe_Total", "Estado", "Fecha_Pedido"])
    orders.append(
        [
            "PO-2026-0703",
            "P001",
            "B98120774",
            None if incomplete else 1802.90,
            "ABIERTO",
            "2026-04-21",
        ]
    )
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
async def payment_api(settings):
    definition = pack.definition()
    async with session_factory() as session:
        process = await rows.process(session, "Payment API " + uuid.uuid4().hex)
        user = User(name="Operator", email=uuid.uuid4().hex + "@test.invalid", role="operator")
        session.add(user)
        await session.flush()
        for spec in definition["decision_types"]:
            session.add(DecisionType(process_id=process.id, **spec))
        for spec in definition["symbols"]:
            session.add(Symbol(process_id=process.id, **spec))
        for rule in pack.rules():
            session.add(
                Rule(
                    process_id=process.id,
                    text=rule.text,
                    type=rule.type,
                    decision=rule.decision,
                    code=rule.code,
                    hash=rule.hash,
                    status="active",
                )
            )
        erp = Source(
            process_id=process.id,
            name="erp",
            origin="test HTTP snapshot",
            rows=[
                {
                    "entry_id": "AS-1",
                    "purchase_order": "PO-2026-0703",
                    "supplier_id": "P001",
                    "nif": "B98120774",
                    "amount": "1802.90",
                    "status": "PENDIENTE",
                }
            ],
        )
        session.add(erp)
        await rows.publish_fixture(session, process.id)
        process_id, user_id, erp_id = process.id, user.id, erp.id
    service = ExtractionService(settings, NoOCR(), NoVLM())
    app.dependency_overrides[current_service] = lambda: service
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
            headers={"X-User-Id": str(user_id)},
        ) as client:
            yield client, process_id, service, erp_id
    finally:
        app.dependency_overrides.pop(current_service)
        await engine.dispose()


async def load_sources(client, process_id, incomplete=False):
    return await client.post(
        f"/processes/{process_id}/sources/workbook",
        files={"file": ("master.xlsx", workbook(incomplete))},
        data={"cut_off_date": "2026-09-19"},
    )


async def test_upload_to_decision_to_export_uses_document_values_and_real_rules(payment_api):
    client, process_id, _, erp_id = payment_api
    sources = await load_sources(client, process_id)
    assert sources.status_code == 201, sources.text
    assert {s["name"] for s in sources.json()["sources"]} == {"suppliers", "orders", "parameters"}
    content = pdf_bytes(VALID)
    upload = await client.post(
        f"/processes/{process_id}/files", files={"file": ("invoice.pdf", content)}
    )
    assert upload.status_code == 201, upload.text
    data = upload.json()
    assert data["status"] == "PENDING"
    assert data["symbols"]["issuer_nif"]["value"] == "B98120774"
    assert data["symbols"]["total"]["value"] == "1802.90"
    assert data["symbols"]["total"]["origin"] == "document:" + data["extraction"]["id"]
    instance_id = data["instance_id"]
    trace = (await client.get(f"/instances/{instance_id}")).json()["events"]
    assert trace[0]["data"]["source_ids"]["erp"] == erp_id
    original = await client.get(f"/instances/{instance_id}/file")
    assert original.content == content
    assert original.headers["content-type"] == "application/pdf"
    visible_sources = (await client.get(f"/processes/{process_id}/sources")).json()
    assert {source["name"] for source in visible_sources} == {
        "erp",
        "suppliers",
        "orders",
        "parameters",
    }
    source_rows = (await client.get(f"/processes/{process_id}/sources/orders")).json()
    assert source_rows["data"][0]["purchase_order"] == "PO-2026-0703"
    process_events = (await client.get(f"/processes/{process_id}/events")).json()
    assert {event["step"] for event in process_events} == {
        "publish_process_version",
        "upload_workbook",
        "load_workbook",
        "upload_document",
        "store_file",
        "extraction",
        "native_text",
        "ingest_document",
    }
    assert (await client.get(f"/processes/{process_id}/export")).status_code == 409
    run = await client.post(f"/processes/{process_id}/run")
    assert run.json() == {"decided": 1, "by_decision": {"PAGAR": 1}}, run.text
    assert (await client.post(f"/processes/{process_id}/run")).json()["decided"] == 0
    export = await client.get(f"/processes/{process_id}/export")
    assert json.loads(export.text) == {"file_id": "invoice.pdf", "result": "PAGAR"}
    detail = (await client.get(f"/instances/{instance_id}")).json()
    assert len(detail["decisions"]) == 1
    assert len(detail["decisions"][0]["results"]) == 17
    summary = (await client.get(f"/processes/{process_id}/summary")).json()
    assert summary["instances"] == 1 and summary["by_decision"] == {"PAGAR": 1}
    assert (await client.post(f"/instances/{instance_id}/extract", json={})).status_code == 409
    assert (await client.get(f"/instances/{instance_id}")).json() == detail


async def test_unverified_identifier_cannot_be_promoted_from_proposal(payment_api):
    from .conftest import lines

    client, process_id, service, _ = payment_api
    await load_sources(client, process_id)
    service.ocr.recognize = lambda *args: lines(VALID, "ocr", 0.99)
    response = await client.post(
        f"/processes/{process_id}/files",
        files={"file": ("scan.pdf", pdf_bytes(""))},
        data={"vlm": "false", "jev": "false"},
    )
    assert response.status_code == 201, response.text
    assert response.json()["extraction"]["fields"]["payment_iban"]["proposed_value"]
    assert response.json()["symbols"]["iban"]["value"] is None
    run = await client.post(f"/processes/{process_id}/run")
    assert run.json()["by_decision"] == {"ESCALAR": 1}, run.text


async def test_scanned_pdf_upload_uses_both_ocr_readers_before_decision(payment_api):
    from .conftest import lines

    client, process_id, service, _ = payment_api
    assert (await load_sources(client, process_id)).status_code == 201
    calls = []

    def read(image, page, size):
        calls.append((image, page, size))
        return lines(VALID, "ocr", 0.99)

    service.ocr.recognize = read
    service.ocr.verify = read
    response = await client.post(
        f"/processes/{process_id}/files",
        files={"file": ("scan.pdf", pdf_bytes(""))},
        data={"ocr": "true", "vlm": "false", "jev": "false"},
    )
    assert response.status_code == 201, response.text
    upload = response.json()
    assert len(calls) == 2
    assert all(image and page == 1 for image, page, _ in calls)
    assert upload["extraction"]["metrics"]["ocr_calls"] == 2
    assert upload["extraction"]["metrics"]["ocr_verification_calls"] == 1
    assert upload["extraction"]["metrics"]["vlm_calls"] == 0
    assert upload["extraction"]["metrics"]["jev_calls"] == 0
    assert upload["symbols"]["issuer_nif"]["value"] == "B98120774"
    assert upload["symbols"]["total"]["value"] == "1802.90"
    run = await client.post(f"/processes/{process_id}/run")
    assert run.status_code == 200, run.text
    assert run.json()["by_decision"] == {"PAGAR": 1}, run.text


async def test_printed_impossible_date_reaches_the_invalid_date_rule(payment_api):
    client, process_id, _, _ = payment_api
    await load_sources(client, process_id)
    upload = await client.post(
        f"/processes/{process_id}/files",
        files={
            "file": ("bad-date.pdf", pdf_bytes(VALID.replace("21/04/2026", "31/02/2026"))),
        },
    )
    assert upload.json()["symbols"]["date"]["value"] == "2026-02-31"
    run = await client.post(f"/processes/{process_id}/run")
    assert run.json()["by_decision"] == {"NO_PAGAR": 1}, run.text


async def test_workbook_validation_is_atomic_and_never_overwrites_erp(payment_api):
    client, process_id, _, erp_id = payment_api
    assert (await load_sources(client, process_id)).status_code == 201
    assert (await load_sources(client, process_id, incomplete=True)).status_code == 422
    async with session_factory() as session:
        snapshots = list(
            await session.scalars(select(Source).where(Source.process_id == process_id))
        )
        assert len(snapshots) == 4
        assert next(s for s in snapshots if s.name == "erp").id == erp_id
    assert (await load_sources(client, process_id)).status_code == 201
    async with session_factory() as session:
        snapshots = list(
            await session.scalars(select(Source).where(Source.process_id == process_id))
        )
        assert len(snapshots) == 7


async def test_pending_reextraction_uses_new_snapshots_and_keeps_both_events(payment_api):
    client, process_id, _, _ = payment_api
    upload = (
        await client.post(
            f"/processes/{process_id}/files",
            files={
                "file": ("invoice.pdf", pdf_bytes(VALID)),
            },
        )
    ).json()
    loaded = await load_sources(client, process_id)
    fresh = await client.post(f"/instances/{upload['instance_id']}/extract", json={})
    assert fresh.status_code == 200, fresh.text
    async with session_factory() as session:
        events = list(
            await session.scalars(
                select(Event)
                .where(
                    Event.instance_id == upload["instance_id"],
                    Event.step.in_(("ingest_document", "extract_document")),
                )
                .order_by(Event.id)
            )
        )
        assert [e.step for e in events] == ["ingest_document", "extract_document"]
        assert events[1].data["source_ids"]["suppliers"] == loaded.json()["sources"][0]["id"]
        instance = await session.get(Instance, upload["instance_id"])
        assert instance.symbols == fresh.json()["symbols"]
    document = await client.get(f"/instances/{upload['instance_id']}/document")
    assert document.json() == fresh.json()["extraction"]
    process_events = (await client.get(f"/processes/{process_id}/events")).json()
    assert "extract_document" in {event["step"] for event in process_events}


async def test_same_source_rows_reuse_evidence_but_changed_schema_refreshes(payment_api):
    client, process_id, service, _ = payment_api
    assert (await load_sources(client, process_id)).status_code == 201
    content = pdf_bytes(VALID + "\n" + uuid.uuid4().hex)
    endpoint = f"/processes/{process_id}/files"
    first = (await client.post(endpoint, files={"file": ("stable.pdf", content)})).json()
    assert (await load_sources(client, process_id)).status_code == 201
    original_extract = service.extract
    service.extract = lambda *_: pytest.fail("Equal source rows must reuse evidence")
    try:
        unchanged = await client.post(endpoint, files={"file": ("stable.pdf", content)})
    finally:
        service.extract = original_extract
    assert unchanged.status_code == 201, unchanged.text
    reused = unchanged.json()["extraction"]
    assert reused["id"] == first["extraction"]["id"]
    assert reused["data"] == first["extraction"]["data"]
    assert reused["cache_hit"] is True
    for reader in ("ocr", "vlm", "jev"):
        assert reused["metrics"][f"{reader}_calls_this_request"] == 0
        assert reused["metrics"][f"{reader}_cache_hits_this_request"] == 0
    document = await client.get(f"/instances/{first['instance_id']}/document")
    assert document.json() == first["extraction"]

    async with session_factory() as session:
        symbol = await session.get(Symbol, (process_id, "total"))
        symbol.description = "Changed extraction contract"
        await rows.publish_fixture(session, process_id)
    refreshed = await client.post(endpoint, files={"file": ("stable.pdf", content)})
    assert refreshed.status_code == 201, refreshed.text
    assert refreshed.json()["extraction"]["id"] != first["extraction"]["id"]
    async with session_factory() as session:
        events = list(
            await session.scalars(
                select(Event)
                .where(
                    Event.instance_id == first["instance_id"],
                    Event.step.in_(("ingest_document", "extract_document")),
                )
                .order_by(Event.id)
            )
        )
    assert [event.step for event in events] == ["ingest_document", "extract_document"]
    async with session_factory() as session:
        await session.delete(await session.get(Symbol, (process_id, "iban")))
        await rows.publish_fixture(session, process_id)
    generic = await client.post(f"/instances/{first['instance_id']}/extract", json={})
    assert generic.status_code == 200, generic.text
    assert generic.json()["symbols"] is None
    async with session_factory() as session:
        instance = await session.get(Instance, first["instance_id"])
        assert instance.symbols is None


async def test_concurrent_runs_do_not_duplicate_decisions(payment_api):
    client, process_id, _, _ = payment_api
    await load_sources(client, process_id)
    upload = await client.post(
        f"/processes/{process_id}/files",
        files={
            "file": ("invoice.pdf", pdf_bytes(VALID)),
        },
    )
    assert upload.status_code == 201
    first, second = await asyncio.gather(
        client.post(f"/processes/{process_id}/run"),
        client.post(f"/processes/{process_id}/run"),
    )
    assert sorted([first.json()["decided"], second.json()["decided"]]) == [0, 1]
    detail = await client.get(f"/instances/{upload.json()['instance_id']}")
    assert len(detail.json()["decisions"]) == 1
