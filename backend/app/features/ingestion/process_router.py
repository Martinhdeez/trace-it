"""Process-facing ingestion routes composed into the main dev application."""

import logging
import mimetypes
import zipfile
from datetime import date
from typing import Annotated, Literal
from urllib.parse import quote
from xml.etree.ElementTree import ParseError

import pymupdf
from fastapi import APIRouter, Depends, File, Form, Response, UploadFile
from fastapi.concurrency import run_in_threadpool
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

from app.common.exceptions import TraceError
from app.core import events
from app.core.database import Session
from app.features.sources.workbook import load_workbook
from app.features.users.dependencies import CurrentUser

from . import process_service
from .errors import InvalidDocumentError
from .extraction_plan import ExtractionPlan, load_extraction_plan
from .pdf.locations import DocumentLocations, locate_document, render_page
from .process_extraction import read_document, reextract_document
from .runtime import current_service
from .schemas import CriticalField, ExtractionResult, ExtractOptions
from .service import ExtractionService

logger = logging.getLogger(__name__)
router = APIRouter(tags=["ingestion"])
Service = Annotated[ExtractionService, Depends(current_service)]


class LoadedSource(BaseModel):
    id: int
    name: str
    rows: int


class WorkbookUpload(BaseModel):
    process_id: int
    file_hash: str
    extraction_id: str
    sources: list[LoadedSource]
    warnings: list[dict]


class DocumentUpload(BaseModel):
    instance_id: int
    process_id: int
    name: str
    file_hash: str
    status: str
    created: bool
    extraction: ExtractionResult
    symbols: dict | None = None


@router.get(
    "/processes/{process_id}/extraction-plan",
    operation_id="getProcessExtractionPlan",
    summary="Inspect the current fields and rule dependencies used for document extraction",
)
async def extraction_plan(process_id: int, session: Session, user: CurrentUser) -> dict:
    await process_service.require_process(session, process_id)
    plan: ExtractionPlan = await load_extraction_plan(session, process_id)
    return {**plan.model_dump(mode="json"), "fingerprint": plan.fingerprint}


@router.post(
    "/processes/{process_id}/files",
    status_code=201,
    operation_id="uploadProcessDocument",
    summary="Store a PDF and its reading evidence as an instance of a process",
    response_model=DocumentUpload,
)
async def upload_document(
    process_id: int,
    session: Session,
    user: CurrentUser,
    service: Service,
    file: Annotated[UploadFile, File()],
    ocr: Annotated[bool, Form()] = True,
    vlm: Annotated[bool | None, Form()] = None,
    jev: Annotated[bool | None, Form()] = None,
    verify_fields: Annotated[list[CriticalField] | None, Form()] = None,
    mode: Annotated[Literal["local", "api", "hybrid"] | None, Form()] = None,
):
    try:
        from app.features.versions.service import lock

        await lock(session, process_id)
        with events.span("upload_document", process_id=process_id, file=file.filename) as span:
            try:
                with events.span("store_file"):
                    item = await run_in_threadpool(service.ingest, file.file, file.filename)
                if item["kind"] != "invoice":
                    raise InvalidDocumentError(
                        "Process documents must be PDF; use /v1/extractions to inspect a workbook"
                    )
                options = ExtractOptions(
                    mode=mode, ocr=ocr, vlm=vlm, jev=jev, verify_fields=verify_fields or []
                )
                from .runtime import for_process

                service, options = await for_process(session, process_id, service, options)
                instance, stored = await process_service.existing_document(
                    session, process_id, item, service, options
                )
                if stored is not None:
                    upload = process_service.stored_upload(instance, stored)
                    span.set(instance_id=instance.id, created=False)
                    return upload
                if instance is not None:
                    upload = await reextract_document(
                        session, instance.id, user.id, service, options
                    )
                    span.set(instance_id=instance.id, created=False)
                    return upload
                result, symbols, context = await read_document(
                    session,
                    process_id,
                    service,
                    item,
                    options,
                )
            except (ValueError, pymupdf.FileDataError) as exc:
                raise InvalidDocumentError("Invalid or unsupported PDF document") from exc
            except TraceError:
                raise
            except Exception as exc:
                logger.exception("Document extraction failed")
                raise TraceError("Document extraction failed; inspect server logs") from exc
            content = await run_in_threadpool((service.objects / item["sha256"]).read_bytes)
            upload = await process_service.attach_document(
                session, process_id, user.id, content, result, symbols, context
            )
            span.set(instance_id=upload["instance_id"], created=upload["created"])
            return upload
    finally:
        await file.close()


@router.get(
    "/instances/{instance_id}/document",
    operation_id="getInstanceDocument",
    summary="Read persisted document evidence from PostgreSQL",
    response_model=ExtractionResult,
)
async def get_document(instance_id: int, session: Session, user: CurrentUser):
    return await process_service.document_result(session, instance_id)


@router.get("/instances/{instance_id}/document/locations", response_model=DocumentLocations)
async def get_document_locations(
    instance_id: int, session: Session, user: CurrentUser, service: Service
):
    from .payment_verification import PAYMENT_FIELDS

    event = await process_service.document_event(session, instance_id)
    result = ExtractionResult.model_validate(event.data["extraction"])
    _, content = await process_service.document_content(session, instance_id)
    locations = await run_in_threadpool(
        locate_document, content, result, service.ocr, service.settings
    )
    locations.symbol_fields = {name: name for name in result.fields}
    # Use the saved adapter, never the process's current (possibly changed) schema.
    if event.data.get("adapter", "").startswith("invoice-payment"):
        locations.symbol_fields.update(
            {symbol: field for symbol, field in PAYMENT_FIELDS.items() if field in result.fields}
        )
    return locations


@router.get("/instances/{instance_id}/document/pages/{page_number}", response_class=Response)
async def get_document_page(
    instance_id: int, page_number: int, session: Session, user: CurrentUser
):
    _, content = await process_service.document_content(session, instance_id)
    image = await run_in_threadpool(render_page, content, page_number)
    return Response(
        image, media_type="image/png", headers={"Cache-Control": "private, max-age=3600"}
    )


@router.post(
    "/instances/{instance_id}/extract",
    operation_id="extractInstanceDocument",
    summary="Re-extract a pending document using current symbols, rules and source snapshots",
    response_model=DocumentUpload,
    responses={409: {"description": "The instance is already decided"}},
)
async def extract_instance(
    instance_id: int,
    session: Session,
    user: CurrentUser,
    service: Service,
    options: ExtractOptions,
):
    return await reextract_document(session, instance_id, user.id, service, options)


@router.post(
    "/processes/{process_id}/sources/workbook",
    status_code=201,
    operation_id="uploadProcessWorkbook",
    response_model=WorkbookUpload,
    summary="Load supplier and order snapshots from an XLSX workbook",
)
async def upload_workbook(
    process_id: int,
    session: Session,
    user: CurrentUser,
    service: Service,
    file: Annotated[UploadFile, File()],
    cut_off_date: Annotated[date | None, Form()] = None,
):
    try:
        await process_service.require_process(session, process_id)
        with events.span("upload_workbook", process_id=process_id, file=file.filename):
            item = await run_in_threadpool(service.ingest, file.file, file.filename)
            if item["kind"] != "workbook":
                raise InvalidDocumentError("Reference sources must be an XLSX workbook")
            result = await run_in_threadpool(
                service.extract, item, ExtractOptions(ocr=False, vlm=False, jev=False)
            )
            content = await run_in_threadpool((service.objects / item["sha256"]).read_bytes)
            return await load_workbook(session, process_id, user.id, result, content, cut_off_date)
    except (ValueError, zipfile.BadZipFile, InvalidFileException, ParseError) as exc:
        raise InvalidDocumentError("Invalid or unsupported workbook") from exc
    finally:
        await file.close()


@router.get(
    "/instances/{instance_id}/file",
    operation_id="getInstanceFile",
    summary="The file the instance was made from, byte for byte (usually a PDF)",
    response_class=Response,
    responses={200: {"content": {"application/pdf": {}}}},
)
async def get_file(instance_id: int, session: Session) -> Response:
    name, content = await process_service.document_content(session, instance_id)
    media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
    # RFC 5987: the name may carry accents; `filename*` keeps them for the browser.
    disposition = f"inline; filename*=UTF-8''{quote(name)}"
    return Response(content, media_type=media_type, headers={"Content-Disposition": disposition})
