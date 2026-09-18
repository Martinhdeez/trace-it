import io
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from app.common.extraction import Candidate, Evidence, ExtractedField
from app.common.normalization import clean_text, excel_decimal, fold, iban, identifier

from .config import WorkbookLimits

ALIASES = {
    "ID": "supplier_ref",
    "PROVEEDORID": "supplier_ref",
    "PROVEEDOR ID": "supplier_ref",
    "SUPPLIER ID": "supplier_ref",
    "PROVEEDOR_ID": "supplier_ref",
    "NIF": "supplier_tax_id",
    "CIF": "supplier_tax_id",
    "TAX ID": "supplier_tax_id",
    "IBAN": "authorized_iban",
    "RAZON SOCIAL": "supplier_name",
    "NOMBRE": "supplier_name",
    "PEDIDO": "purchase_order_ref",
    "PO": "purchase_order_ref",
    "ORDER ID": "purchase_order_ref",
    "IMPORTE_TOTAL": "expected_gross_amount",
    "IMPORTE TOTAL": "expected_gross_amount",
    "TOTAL": "expected_gross_amount",
    "AMOUNT": "expected_gross_amount",
    "ESTADO": "state",
    "STATUS": "state",
    "FECHA_PEDIDO": "ordered_on",
    "FECHA PEDIDO": "ordered_on",
    "MONEDA": "currency",
    "CURRENCY": "currency",
    "ASIENTO_ID": "ledger_entry_id",
    "FECHA_REGISTRO": "recorded_on",
    "IMPORTE_ESPERADO": "expected_gross_amount",
}


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def extract_workbook(content: bytes, settings: WorkbookLimits):
    numeric_xml = {}
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(f.file_size for f in members) > 200 * 1024 * 1024:
                raise ValueError("Workbook expanded size exceeds limit")
            if any(f.flag_bits & 1 for f in members):
                raise ValueError("Encrypted workbook is unsupported")
            # Retain numeric lexical values before openpyxl converts them to binary
            # floats, and check real coordinates even if <dimension> is misleading.
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for member in members:
                if not member.filename.startswith("xl/worksheets/") or not member.filename.endswith(
                    ".xml"
                ):
                    continue
                cells = {}
                with archive.open(member) as stream:
                    for _, element in ET.iterparse(stream, events=("end",)):
                        if element.tag == ns + "c":
                            coordinate = element.get("r")
                            if coordinate:
                                r, c = coordinate_to_tuple(coordinate)
                                if r > settings.max_excel_rows or c > settings.max_excel_cols:
                                    raise ValueError("Sheet dimensions exceed limit")
                                if element.get("t", "n") == "n":
                                    cells[coordinate] = element.findtext(ns + "v")
                            element.clear()
                        elif element.tag == ns + "row":
                            element.clear()
                numeric_xml[member.filename] = cells
    except zipfile.BadZipFile as exc:
        raise ValueError("Expected an XLSX workbook") from exc
    try:
        formulas = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (KeyError, OSError) as exc:
        raise ValueError("Missing or invalid XLSX workbook parts") from exc
    try:
        values = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except Exception:
        formulas.close()
        raise
    warnings, sheets, records = [], [], []
    seen = {}
    total_cells = 0
    try:
        for sheet in formulas:
            if (sheet.max_row or 0) > settings.max_excel_rows or (
                sheet.max_column or 0
            ) > settings.max_excel_cols:
                raise ValueError(f"Sheet dimensions exceed limit: {sheet.title}")
            sheet.reset_dimensions()
            values[sheet.title].reset_dimensions()
            lexical_values = numeric_xml.get(sheet._worksheet_path, {})
            rows = []
            cached_rows = values[sheet.title].iter_rows()
            for row_number, (row, cached_row) in enumerate(
                zip(sheet.iter_rows(), cached_rows, strict=True), 1
            ):
                total_cells += len(row)
                if row_number > settings.max_excel_rows or len(row) > settings.max_excel_cols:
                    raise ValueError(f"Sheet dimensions exceed limit: {sheet.title}")
                if total_cells > settings.max_excel_cells:
                    raise ValueError("Workbook cell count exceeds limit")
                cells = []
                for cell, cached in zip(row, cached_row, strict=True):
                    if cell.value is None:
                        continue
                    raw, val = cell.value, cell.value
                    if cell.data_type == "f":
                        val = cached.value
                        warnings.append(
                            {
                                "code": "FORMULA_CACHED_VALUE"
                                if val is not None
                                else "FORMULA_UNEVALUATED",
                                "locator": f"{sheet.title}!{cell.coordinate}",
                            }
                        )
                    if isinstance(val, str):
                        val = clean_text(val)
                    cells.append(
                        {
                            "cell": cell.coordinate,
                            "column": cell.column,
                            "raw": json_value(raw),
                            "value": json_value(val),
                            "formula": cell.data_type == "f",
                            "excel_type": cell.data_type,
                            "xml_numeric_value": lexical_values.get(cell.coordinate),
                            "number_format": cell.number_format,
                        }
                    )
                if cells:
                    rows.append({"row": row_number, "cells": cells})
            header = None
            mapping = {}
            role = "unstructured"
            for row in rows[:25]:
                matched = {
                    c["column"]: ALIASES[fold(str(c["value"]))]
                    for c in row["cells"]
                    if fold(str(c["value"])) in ALIASES
                }
                names = set(matched.values())
                if "ledger_entry_id" in names and "purchase_order_ref" in names:
                    role = "ledger"
                elif "purchase_order_ref" in names and "expected_gross_amount" in names:
                    role = "purchase_orders"
                elif {"supplier_tax_id", "authorized_iban"} <= names:
                    role = "suppliers"
                else:
                    continue
                header, mapping = row["row"], matched
                break
            if "NORMA" in fold(sheet.title) or "POLICY" in fold(sheet.title):
                role = "payment_policy"
            sheets.append(
                {
                    "name": sheet.title,
                    "visibility": sheet.sheet_state,
                    "role": role,
                    "header_row": header,
                    "rows": rows,
                }
            )
            if not header:
                if rows and role == "unstructured":
                    warnings.append({"code": "UNRECOGNIZED_SCHEMA", "sheet": sheet.title})
                continue
            for row in rows:
                if row["row"] <= header:
                    continue
                cells = {c["column"]: c for c in row["cells"]}
                fields = {}
                for col, name in mapping.items():
                    cell = cells.get(col)
                    field = ExtractedField()
                    if cell and cell["value"] not in (None, ""):
                        raw = str(cell["value"])
                        status, value, error = "OBSERVED", raw, None
                        try:
                            if name == "expected_gross_amount":
                                value = excel_decimal(cell["value"])
                            elif name == "authorized_iban":
                                value = iban(raw)
                            elif name == "supplier_tax_id":
                                value = identifier(raw)
                            elif name in {
                                "supplier_ref",
                                "purchase_order_ref",
                                "state",
                                "currency",
                            }:
                                value = clean_text(raw).upper()
                        except ValueError as exc:
                            value, status, error = None, "INVALID", str(exc)
                        if cell["excel_type"] == "e":
                            value, status, error = None, "INVALID", "Excel error cell"
                        if cell["formula"]:
                            status = "UNVERIFIED"
                        candidate = Candidate(
                            value=value,
                            raw=str(cell["raw"]),
                            error=error,
                            evidence=Evidence(
                                locator=f"{sheet.title}!{cell['cell']}",
                                text=str(cell["raw"]),
                                method="xlsx",
                            ),
                        )
                        field = ExtractedField(
                            value=value,
                            status=status,
                            candidates=[candidate],
                            origin="AUTHORITATIVE_SOURCE",
                        )
                    fields.setdefault(name, []).append(field)
                flattened = {}
                for key, alternatives in fields.items():
                    populated = [f for f in alternatives if f.status != "MISSING"]
                    field = populated[0] if populated else alternatives[0]
                    if len({f.value for f in populated}) > 1:
                        field = ExtractedField(
                            status="AMBIGUOUS",
                            candidates=[c for f in populated for c in f.candidates],
                        )
                    flattened[key] = field.model_dump()
                required = {
                    "suppliers": ("supplier_tax_id", "authorized_iban"),
                    "purchase_orders": ("purchase_order_ref", "expected_gross_amount", "state"),
                    "ledger": ("ledger_entry_id", "purchase_order_ref", "state"),
                }.get(role, ())
                for key in required:
                    flattened.setdefault(key, ExtractedField().model_dump())
                if role in {"purchase_orders", "ledger"} and not any(
                    flattened.get(key, {}).get("value")
                    for key in ("supplier_ref", "supplier_tax_id")
                ):
                    flattened.setdefault("supplier_ref", ExtractedField().model_dump())
                record = {
                    "role": role,
                    "sheet": sheet.title,
                    "row": row["row"],
                    "fields": flattened,
                }
                primary = (
                    "purchase_order_ref"
                    if role == "purchase_orders"
                    else ("ledger_entry_id" if role == "ledger" else "supplier_ref")
                )
                identity = flattened.get(primary, {}).get("value") or flattened.get(
                    "supplier_tax_id", {}
                ).get("value")
                key = role, identity
                signature = tuple(sorted((k, v["value"]) for k, v in flattened.items()))
                if identity and key in seen:
                    prior, previous_signature = seen[key]
                    warnings.append(
                        {
                            "code": "DUPLICATE_IDENTICAL"
                            if signature == previous_signature
                            else "DUPLICATE_CONFLICT",
                            "identity": identity,
                            "first": prior,
                            "next": f"{sheet.title}!{row['row']}",
                        }
                    )
                elif identity:
                    seen[key] = (f"{sheet.title}!{row['row']}", signature)
                for key, value in flattened.items():
                    if value["status"] != "OBSERVED":
                        warnings.append(
                            {
                                "code": "FIELD_" + value["status"],
                                "field": key,
                                "locator": f"{sheet.title}!{row['row']}",
                            }
                        )
                records.append(record)
    finally:
        formulas.close()
        values.close()
    return {"sheets": sheets, "records": records, "cells_scanned": total_cells}, warnings
