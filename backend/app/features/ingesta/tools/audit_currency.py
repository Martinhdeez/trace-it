"""Cross-check missing PDF currency against independent text extraction and workbook context."""

import argparse
import hashlib
import json
import re
import shutil
import subprocess
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from app.features.ingesta.config import Settings
from app.features.ingesta.pdf.invoice import CURRENCY_CODE, parse_invoice
from app.features.ingesta.pdf.native import native_pages

MARKERS = re.compile(CURRENCY_CODE + r"|\bEUROS?\b|[€$£¥]", re.I)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, default=Path("reports/currency-audit"))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    settings = Settings()
    workbook_path = args.input / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"
    book = load_workbook(workbook_path, data_only=True)
    orders = {
        row[0].value: {
            "row": row[0].row,
            "supplier_ref": row[1].value,
            "tax_id": row[2].value,
            "gross": str(row[3].value),
        }
        for row in list(book["Pedidos_2026"].iter_rows())[1:]
        if row[0].value
    }
    workbook_markers, formats = [], Counter()
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    formats[cell.number_format] += 1
                    if isinstance(cell.value, str) and MARKERS.search(cell.value):
                        workbook_markers.append(
                            {"locator": f"{sheet.title}!{cell.coordinate}", "text": cell.value}
                        )
    book.close()
    native, missing, observed, supplier_support = [], [], Counter(), defaultdict(list)
    for path in sorted((args.input / "facturas").glob("*.pdf")):
        content = path.read_bytes()
        pages = native_pages(content, settings)
        lines = [line for page in pages for line in page["lines"]]
        if not lines:
            continue
        fields, _ = parse_invoice(lines)
        currency = fields["currency"]
        item = {
            "file_id": path.name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "pages": len(pages),
            "native_markers": MARKERS.findall("\n".join(line.text for line in lines)),
            "first_line": lines[0].text,
            "supplier_tax_id": fields["supplier_tax_id"].value,
            "purchase_order_ref": fields["purchase_order_ref"].value,
            "gross_amount": fields["gross_amount"].value,
            "currency": currency.model_dump(),
        }
        native.append(item)
        if currency.status == "OBSERVED":
            observed[currency.value] += 1
            supplier_support[item["supplier_tax_id"]].append(
                {"file_id": path.name, "currency": currency.value}
            )
        if currency.status == "MISSING":
            with pymupdf.open(path) as doc:
                item["embedded_images"] = sum(len(p.get_images()) for p in doc)
                item["fonts"] = sorted({font[3] for p in doc for font in p.get_fonts()})
            order = orders.get(item["purchase_order_ref"])
            item["matched_order"] = order
            item["order_total_matches"] = bool(
                order
                and item["gross_amount"]
                and abs(Decimal(order["gross"]) - Decimal(item["gross_amount"])) <= Decimal("0.01")
            )
            missing.append(item)

    executable = shutil.which("pdftotext")
    if not executable:
        raise RuntimeError("Install pdftotext for the independent extraction check")

    def check(item):
        path = args.input / "facturas" / item["file_id"]
        result = subprocess.run(
            [executable, "-enc", "UTF-8", str(path), "-"],
            capture_output=True,
            check=True,
            timeout=20,
        )
        text = result.stdout.decode("utf-8")
        item["pdftotext_markers"] = MARKERS.findall(text)
        item["same_supplier_explicit_currencies"] = dict(
            Counter(entry["currency"] for entry in supplier_support[item["supplier_tax_id"]])
        )
        item["same_supplier_example"] = next(iter(supplier_support[item["supplier_tax_id"]]), None)
        return item

    with ThreadPoolExecutor(max_workers=4) as pool:
        missing = list(pool.map(check, missing))
    summary = {
        "native_pdfs": len(native),
        "observed_currencies": dict(observed),
        "missing_currency": len(missing),
        "missing_with_native_markers": sum(bool(x["native_markers"]) for x in missing),
        "missing_with_pdftotext_markers": sum(bool(x["pdftotext_markers"]) for x in missing),
        "missing_with_embedded_images": sum(x["embedded_images"] > 0 for x in missing),
        "headers": dict(Counter(x["first_line"] for x in missing)),
        "template_families": dict(
            Counter(
                "simplified"
                if x["first_line"].startswith("FACTURA SIMPLIFICADA")
                else "standard"
                if x["first_line"] == "FACTURA"
                else "uppercase_supplier_header"
                for x in missing
            )
        ),
        "fonts": dict(Counter(f for x in missing for f in x["fonts"])),
        "pages": dict(Counter(x["pages"] for x in missing)),
        "matched_orders": sum(bool(x["matched_order"]) for x in missing),
        "matching_order_totals": sum(x["order_total_matches"] for x in missing),
        "same_supplier_eur_evidence": sum(
            "EUR" in x["same_supplier_explicit_currencies"] for x in missing
        ),
        "workbook_currency_mentions": workbook_markers,
        "workbook_number_formats": dict(formats),
        "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "interpretation": (
            "Use policy and matching-document evidence to infer batch "
            "currency; numeric matches alone do not establish currency."
        ),
    }
    (args.output / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    with (args.output / "files.jsonl").open("w", encoding="utf-8") as stream:
        for item in missing:
            stream.write(json.dumps(item, ensure_ascii=False) + "\n")
    print(json.dumps(summary, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
