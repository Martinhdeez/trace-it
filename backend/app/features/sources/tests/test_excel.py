import io
from dataclasses import replace

import pytest
from openpyxl import Workbook

from app.features.sources.excel import extract_workbook


@pytest.mark.parametrize("dimension", ["", '<dimension ref="A1:A1"/>'])
def test_workbook_sparse_rows_and_missing_or_false_dimensions(settings, dimension):
    import re
    import zipfile

    book = Workbook()
    sheet = book.active
    sheet.append([None, "PEDIDO", "IMPORTE TOTAL", "ESTADO", "PROVEEDOR ID"])
    sheet.append([None, "PO-2026-0001", 9299.62, "ABIERTO", "P001"])
    original, modified = io.BytesIO(), io.BytesIO()
    book.save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = re.sub(r"<dimension[^>]*/>", dimension, content.decode())
                content = text.replace("<v>9299.62</v>", "<v>9299.620000000001</v>").encode()
            target.writestr(name, content)
    data, _ = extract_workbook(modified.getvalue(), settings)
    assert len(data["records"]) == 1
    assert data["records"][0]["fields"]["expected_gross_amount"]["value"] == "9299.62"
    cell = next(c for c in data["sheets"][0]["rows"][1]["cells"] if c["cell"] == "C2")
    assert cell["xml_numeric_value"] == "9299.620000000001"
    assert cell["number_format"] == "General"


def test_workbook_preserves_formulas_duplicates_and_missing(settings):
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "NIF", "IBAN"])
    sheet.append(["P007", " J40112358 ", "ES55 3159 0012 3487 6512 3407"])
    sheet.append(["P007", "J40112358", "ES55 3159 0012 3487 6512 3407"])
    sheet.append(["P008", None, "=SUM(A1:A2)"])
    stream = io.BytesIO()
    book.save(stream)
    data, warnings = extract_workbook(stream.getvalue(), settings)
    assert len(data["records"]) == 3
    assert data["records"][0]["fields"]["supplier_tax_id"]["value"] == "J40112358"
    assert {w["code"] for w in warnings} >= {
        "DUPLICATE_IDENTICAL",
        "FORMULA_UNEVALUATED",
        "FIELD_MISSING",
    }


def test_workbook_conflicting_headers_and_limits(settings):
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "NIF", "CIF", "IBAN"])
    sheet.append(["P001", "B12345678", "B87654321", "ES4414650100951704302211"])
    stream = io.BytesIO()
    book.save(stream)
    data, warnings = extract_workbook(stream.getvalue(), settings)
    assert data["records"][0]["fields"]["supplier_tax_id"]["status"] == "AMBIGUOUS"
    assert any(w["code"] == "FIELD_AMBIGUOUS" for w in warnings)
    with pytest.raises(ValueError, match="dimensions"):
        extract_workbook(stream.getvalue(), replace(settings, max_excel_cols=3))


def test_order_sheet_missing_required_column_is_not_complete(settings):
    book = Workbook()
    book.active.append(["PEDIDO", "IMPORTE TOTAL"])
    book.active.append(["PO-2026-0001", 100])
    stream = io.BytesIO()
    book.save(stream)
    data, warnings = extract_workbook(stream.getvalue(), settings)
    fields = data["records"][0]["fields"]
    assert fields["state"]["status"] == "MISSING"
    assert fields["supplier_ref"]["status"] == "MISSING"
    assert {w["field"] for w in warnings if w["code"] == "FIELD_MISSING"} == {
        "state",
        "supplier_ref",
    }
