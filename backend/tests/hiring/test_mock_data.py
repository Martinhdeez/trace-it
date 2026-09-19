"""The committed hiring mock data is exactly what `tools/hiring_mock.py` writes, and its
answer key agrees with the policy written as code. No LLM, no database."""

import json
import runpy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "processes" / "hiring-screening" / "data"
MOCK = runpy.run_path(str(ROOT / "tools" / "hiring_mock.py"))


def committed_expected() -> list[dict]:
    lines = (DATA / "expected.jsonl").read_text(encoding="utf-8").splitlines()
    return [json.loads(line) for line in lines if line.strip()]


def test_committed_files_match_the_generator():
    candidates, tables, expected = MOCK["generate"]()
    assert committed_expected() == expected
    for candidate in candidates:
        assert (DATA / "cvs" / candidate.file_id).read_bytes() == MOCK["render"](candidate), (
            candidate.file_id
        )
    assert sorted(p.name for p in (DATA / "cvs").glob("*.pdf")) == [c.file_id for c in candidates]
    book = openpyxl.load_workbook(DATA / "hiring-reference.xlsx", read_only=True)
    assert book.sheetnames == list(tables)
    for name, rows in tables.items():
        cells = list(book[name].iter_rows(values_only=True))
        assert cells[0] == tuple(rows[0])
        # An empty cell reads back as None; the generator writes "".
        read = [dict(zip(cells[0], row, strict=True)) for row in cells[1:]]
        assert [{k: v if v is not None else "" for k, v in r.items()} for r in read] == rows


def test_answer_key_follows_the_policy_as_code():
    _, tables, _ = MOCK["generate"]()
    rows = committed_expected()
    for row in rows:
        others = [r["symbols"] for r in rows if r["file_id"] != row["file_id"]]
        decision, why = MOCK["outcome"](row["symbols"], tables, others)
        assert (decision, why) == (row["expected"], row["why"]), row["file_id"]


def test_every_trap_and_every_outcome_is_covered():
    rows = committed_expected()
    assert {row["expected"] for row in rows} == {"INTERVIEW", "REJECT", "REVIEW"}
    categories = {row["category"] for row in rows}
    assert categories == {name for name, _, _ in MOCK["CATEGORIES"]}
    assert sum(row["scanned"] for row in rows) == 2
    assert sum(row["language"] == "es" for row in rows) == 2
    # Precedence: a duplicate pair shares an email and both go to a person.
    duplicates = [row for row in rows if row["category"] == "DUPLICATE_CANDIDATE"]
    assert len({row["symbols"]["email"] for row in duplicates}) == 1
    assert all(row["expected"] == "REVIEW" for row in duplicates)
